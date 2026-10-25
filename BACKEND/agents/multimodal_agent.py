import io
import os
import time
import base64

from google import genai
from google.genai import types
from PIL import Image
import pandas as pd
import docx
from dotenv import load_dotenv
from utils.logger import get_logger

load_dotenv()
logger = get_logger(__name__)

VISION_MODELS = [m.strip() for m in
                 os.getenv("VISION_MODELS", "gemini-flash-latest,gemini-flash-lite-latest").split(",") if m.strip()]
RETRIES = 3
WAITS = (5, 12)                                        
GEMINI_TIMEOUT_S = int(os.getenv("VISION_TIMEOUT_S", "40"))   
TOTAL_BUDGET_S = int(os.getenv("VISION_BUDGET_S", "90"))      
FALLBACK_MAX_CHARS = 12000                              
TRANSIENT = ("503", "UNAVAILABLE", "429", "RESOURCE_EXHAUSTED", "500", "504", "DEADLINE", "overloaded",
             "timed out", "timeout", "Timeout")

SUPPORTED_TYPES = {
    "image": [".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"],
    "pdf":   [".pdf"],
    "word":  [".docx", ".doc"],
    "excel": [".xlsx", ".xls"],
    "csv":   [".csv"],
    "text":  [".txt", ".md"],
}


def get_file_type(filename: str) -> str:
    ext = os.path.splitext(filename)[1].lower()
    for ftype, exts in SUPPORTED_TYPES.items():
        if ext in exts:
            return ftype
    return "unknown"


def _is_transient(e: Exception) -> bool:
    msg = str(e)
    return any(k in msg for k in TRANSIENT)


def _friendly(e: Exception) -> str:
    if _is_transient(e):
        return ("Le service d'analyse de fichiers (Gemini) est surchargé ou limité en ce moment. "
                "Réessayez dans quelques instants.")
    return str(e)


def _groq_text(prompt: str) -> str:
    """Envoie un prompt texte a Groq. Necessite : pip install groq"""
    key = os.getenv("GROQ_API_KEY")
    if not key:
        raise RuntimeError("GROQ_API_KEY manquante")
    from groq import Groq
    model = os.getenv("GROQ_MODEL", "openai/gpt-oss-120b")
    r = Groq(api_key=key, timeout=40).chat.completions.create(
        model=model,
        messages=[{"role": "user", "content": prompt}],
        temperature=0.1,
    )
    text = (r.choices[0].message.content or "").strip()
    if not text:
        raise RuntimeError("Reponse vide de Groq")
    return text if text.startswith("[") else "[TEXT] " + text


def _pdf_text(pdf_bytes: bytes) -> str:
    """Extrait le texte d'un PDF en local (pas d'images ni de mise en forme)."""
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(pdf_bytes))
    parts = []
    for i, page in enumerate(reader.pages, 1):
        t = (page.extract_text() or "").strip()
        if t:
            parts.append(f"[Page {i}]\n{t}")
    return "\n\n".join(parts)


BASE_PROMPT = (
    "Tu es un assistant francophone expert en analyse de documents.\n"
    "Reponds TOUJOURS en francais, clairement et structure.\n\n"
    "IMPORTANT — Choisis le format de sortie selon le contenu trouve :\n"
    "- Si le document contient principalement du texte → commence par [TEXT]\n"
    "- Si le document contient un tableau de donnees  → commence par [TABLE] puis JSON [/TABLE]\n"
    "- Si le document contient un graphe ou schema    → commence par [GRAPH] puis JSON [/GRAPH]\n"
    "- Si le document est une image a decrire         → commence par [TEXT]\n\n"
)


class ImageAgent:
    def __init__(self):
        self.client = genai.Client(
            api_key=os.getenv("GEMINI_API_KEY"),
            http_options=types.HttpOptions(timeout=GEMINI_TIMEOUT_S * 1000),  
        )

    def _generate(self, contents) -> str:
        last = None
        start = time.time()
        for model in VISION_MODELS:
            for attempt in range(RETRIES):
                if time.time() - start > TOTAL_BUDGET_S:
                    logger.warning(f"Budget de temps Gemini depasse ({TOTAL_BUDGET_S}s) -> abandon")
                    raise last if last else TimeoutError("Gemini : budget de temps depasse")
                try:
                    text = self.client.models.generate_content(model=model, contents=contents).text
                    if not text:
                        raise RuntimeError("Reponse vide du modele de vision")
                    return text
                except Exception as e:
                    last = e
                    msg = str(e)
                    if "404" in msg or "NOT_FOUND" in msg:
                        logger.warning(f"Modele {model} introuvable -> modele suivant")
                        break                                 
                    if not _is_transient(e):
                        raise                                  
                    if "PerDay" in msg:
                        logger.warning(f"Quota journalier atteint pour {model} -> modele suivant")
                        break                                  
                    logger.warning(f"Gemini {model} indisponible (essai {attempt + 1}/{RETRIES}) : {msg[:90]}")
                    if attempt < RETRIES - 1:
                        time.sleep(WAITS[attempt])
        raise last if last else RuntimeError("Aucun modele de vision configure (VISION_MODELS)")

    def _generate_with_fallback(self, contents, text_prompt: str) -> str:
        """Gemini d'abord ; si indisponible, meme demande envoyee a Groq (texte uniquement)."""
        try:
            return self._generate(contents)
        except Exception as e:
            if not _is_transient(e):
                raise
            logger.warning("Gemini indisponible -> plan B : Groq")
            try:
                return _groq_text(text_prompt[:FALLBACK_MAX_CHARS])
            except Exception as e2:
                logger.warning(f"Plan B Groq indisponible : {str(e2)[:120]}")
                raise e                                      

    def run(self, state: dict) -> dict:
        query    = state.get("query", "Analyse ce fichier en detail.")
        uploaded = state.get("image")

        if uploaded is None:
            logger.warning("ImageAgent appele sans fichier.")
            return {"context": ""}

        filename = uploaded.name if hasattr(uploaded, "name") else "fichier"
        try:
            file_type = get_file_type(filename)
            logger.info(f"Fichier : {filename} | Type : {file_type}")

            if file_type == "image":
                context = self._handle_image(uploaded, query)
            elif file_type == "pdf":
                context = self._handle_pdf_gemini(uploaded, query)
            elif file_type == "word":
                context = self._handle_word(uploaded, query)
            elif file_type in ("excel", "csv"):
                context = self._handle_table(uploaded, query, file_type)
            elif file_type == "text":
                context = self._handle_text(uploaded, query)
            else:
                context = f"[TEXT] Format non supporte : {filename}"

            return {"context": context, "sources": [{"type": "file", "title": filename}]}

        except Exception as e:
            logger.error(f"Erreur ImageAgent : {e}")
            return {"context": f"[TEXT] Impossible d'analyser le fichier : {_friendly(e)}", "sources": []}

    def _handle_image(self, uploaded, query: str) -> str:
        pil_image = Image.open(uploaded)
        prompt = (
            BASE_PROMPT
            + f"Question : {query}\n\n"
            + "Decris l'image : contenu, texte visible, tableaux, graphes, contexte.\n"
            + "REGLES : decris uniquement ce qui est visible. Si c'est une image medicale, n'ecris JAMAIS "
            + "que l'examen est normal ou anormal, qu'il n'y a pas d'anomalie ou de lesion, et ne propose "
            + "aucun diagnostic. Termine par : 'Seul un professionnel de sante peut interpreter cet examen.'"
        )
        
        return self._generate([prompt, pil_image])

    def _handle_pdf_gemini(self, uploaded, query: str) -> str:
        pdf_bytes = uploaded.read()
        uploaded.seek(0)
        pdf_b64 = base64.standard_b64encode(pdf_bytes).decode("utf-8")

        prompt = (
            BASE_PROMPT
            + f"Question : {query}\n\n"
            + "Analyse ce document PDF en entier :\n"
            + "- Lis tout le texte\n"
            + "- Decris les images trouvees\n"
            + "- Extrait les tableaux si presents\n"
            + "- Decris les graphes ou schemas si presents\n"
            + "Reponds a la question de l'utilisateur."
        )
        try:
            return self._generate([{
                "parts": [
                    {"inline_data": {"mime_type": "application/pdf", "data": pdf_b64}},
                    {"text": prompt},
                ]
            }])
        except Exception as e:
            if not _is_transient(e):
                raise
            try:
                local = _pdf_text(pdf_bytes)
            except Exception as e_local:
                logger.warning(f"Extraction locale du PDF impossible : {e_local}")
                raise e
            if len(local.strip()) < 100:
                raise e                                         
            logger.warning("Gemini indisponible -> plan B : texte du PDF extrait en local")
            text_prompt = (
                BASE_PROMPT
                + "Voici le texte extrait d'un PDF (images, tableaux et mise en forme non disponibles) :\n\n"
                + f"{local[:FALLBACK_MAX_CHARS]}\n\n"
                + f"Question : {query}\n\n"
                + "Reponds a la question en te basant uniquement sur ce texte."
            )
            try:
                return _groq_text(text_prompt)
            except Exception as e2:
                logger.warning(f"Plan B Groq indisponible : {str(e2)[:120]}")
                return ("[TEXT] L'analyse automatique est indisponible pour le moment. "
                        "Voici le texte extrait du PDF :\n\n" + local[:3000])

    def _handle_word(self, uploaded, query: str) -> str:
        doc = docx.Document(uploaded)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        text = "\n".join(paragraphs)

        tables_text = ""
        for i, table in enumerate(doc.tables):
            tables_text += f"\nTableau {i+1} :\n"
            for row in table.rows:
                tables_text += " | ".join(cell.text.strip() for cell in row.cells) + "\n"

        full_content = ""
        if text:
            full_content += f"TEXTE DU DOCUMENT :\n{text[:6000]}\n\n"
        if tables_text:
            full_content += f"TABLEAUX TROUVES :\n{tables_text[:2000]}\n\n"
        if not full_content.strip():
            return "[TEXT] Le document Word ne contient pas de contenu extractible."

        prompt = (
            BASE_PROMPT
            + f"Voici le contenu d'un document Word :\n\n{full_content}\n\n"
            + f"Question : {query}\n\n"
            + "Analyse et reponds a la question."
        )
        return self._generate_with_fallback([prompt], prompt)

    def _handle_table(self, uploaded, query: str, file_type: str) -> str:
        df = pd.read_csv(uploaded) if file_type == "csv" else pd.read_excel(uploaded)

        shape, columns = df.shape, list(df.columns)
        preview = df.head(20).to_string(index=False)
        stats = df.describe(include="all").to_string()

        prompt = (
            "Tu es un assistant francophone expert en analyse de donnees.\n"
            "Reponds TOUJOURS en francais, clairement et structure.\n\n"
            f"Tableau de donnees ({shape[0]} lignes, {shape[1]} colonnes) :\n"
            f"Colonnes : {columns}\n\n"
            f"Apercu (20 premieres lignes) :\n{preview}\n\n"
            f"Statistiques :\n{stats}\n\n"
            f"Question : {query}\n\n"
            "Analyse ce tableau et reponds a la question.\n"
            "Reponds en texte simple, sans balises."
        )
        try:
            analysis = self._generate_with_fallback([prompt], prompt)
            if analysis.startswith("[TEXT] "):                 
                analysis = analysis[len("[TEXT] "):]
        except Exception as e:
            
            logger.warning(f"Analyse du tableau indisponible : {e}")
            analysis = f"Analyse automatique indisponible ({_friendly(e)}). Voici les données du fichier."

        return f"[TABLEAU]{df.to_json(orient='records')}\n[ANALYSE]{analysis}"

    def _handle_text(self, uploaded, query: str) -> str:
        text = uploaded.read().decode("utf-8", errors="ignore")[:8000]
        prompt = (
            BASE_PROMPT
            + f"Voici un document texte :\n\n{text}\n\n"
            + f"Question : {query}\n\n"
            + "Analyse et reponds a la question."
        )
        return self._generate_with_fallback([prompt], prompt)